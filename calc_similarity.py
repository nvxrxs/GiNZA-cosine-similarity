import numpy as np
import openpyxl
from scipy.stats import linregress
from sentence_transformers import SentenceTransformer, util

# 1. 日本語埋め込みモデルのロード（初回実行時に自動ダウンロード）
model = SentenceTransformer("pkshatech/GLuCoSE-base-ja")

# 2. スプレッドシート（.xlsx）の読み込み
file_path = "quiz_data.xlsx"
wb = openpyxl.load_workbook(file_path)
sheet = wb["問題"]

# 3. 類似度および傾きの計算処理
# 行3から開始（奇数行: 問題文, 偶数行: ANS）
max_row = sheet.max_row

for row_idx in range(3, max_row, 2):
    ans_cell = sheet.cell(row=row_idx + 1, column=1).value  # A列の偶数行 (ANS)
    if not ans_cell:
        continue
        
    ans_text = str(ans_cell).strip()
    ans_emb = model.encode(ans_text, convert_to_tensor=True)
    
    similarities = []
    col_indices = [2, 3, 4, 5]  # B, C, D, E列
    valid_steps = []
    
    for step_num, col_idx in enumerate(col_indices, start=1):
        segment_cell = sheet.cell(row=row_idx, column=col_idx).value
        
        if segment_cell and str(segment_cell).strip():
            segment_text = str(segment_cell).strip()
            seg_emb = model.encode(segment_text, convert_to_tensor=True)
            
            # コサイン類似度計算
            cos_sim = float(util.cos_sim(seg_emb, ans_emb)[0][0])
            similarities.append(cos_sim)
            valid_steps.append(step_num)
            
            # 偶数行（ANSの行）の該当列にコサイン類似度を書き込み
            sheet.cell(row=row_idx + 1, column=col_idx, value=round(cos_sim, 4))
        else:
            # 空白区間
            sheet.cell(row=row_idx + 1, column=col_idx, value=None)
            
    # 回帰直線の傾きの計算（2区間以上存在する場合）
    if len(similarities) >= 2:
        slope, intercept, r_value, p_value, std_err = linregress(valid_steps, similarities)
        # F列の偶数行に傾きを出力
        sheet.cell(row=row_idx + 1, column=6, value=round(slope, 4))

# 4. 結果を別名で保存
wb.save("quiz_data_analyzed.xlsx")
print("解析が完了し、quiz_data_analyzed.xlsx に保存されました。")